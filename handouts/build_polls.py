#!/usr/bin/env python3
"""Poll question bank for the faculty workshop, 23 September 2026.

    python3 build_polls.py

writes  poll-bank.html   facilitator sheet: questions, answers, one-line explanations
        poll-bank.xlsx   one row per question, for pasting into Socrative's import template

Every number in these questions is one already on a slide in the decks; nothing new is introduced.
No em dashes: house rule.
"""
import html, pathlib
HERE = pathlib.Path(__file__).parent
E = html.escape

# (time, session, [ (question, [options], correct_index or None for opinion, explanation) ])
BANK = [
("Session 1", "Where economic evidence enters health policy", [
 ("Before today, how many published economic evaluations in your own field have you read closely?",
  ["None", "One", "Two to five", "More than five"], None,
  "Baseline only. Re-ask at the close in a different form (the confidence question)."),
 ("Who starts an HTAIn assessment?",
  ["The HTAIn Secretariat, on its own initiative", "A programme, ministry department or state that nominates a topic",
   "The manufacturer of the technology", "A Regional Resource Centre that wants to publish"], 1,
  "The user nominates; the recommendation returns to the user, who decides. HTAIn advises, it does not command."),
]),
("Session 2", "What economic evaluation is, and is not", [
 ("A study reports the cost of treating cervical cancer by stage, with no comparison of options. What is it?",
  ["A full economic evaluation", "A cost description", "A cost-effectiveness analysis", "A cost-minimisation analysis"], 1,
  "Costs only, no alternatives compared: Drummond's cost description box. Useful, but it cannot say what to fund."),
 ("A paper titled \"Cost-effectiveness of X versus Y\" compares only the costs of X and Y. What is it really?",
  ["A cost-effectiveness analysis, as titled", "A cost-utility analysis", "A cost analysis, a partial evaluation", "A cost-benefit analysis"], 2,
  "Two alternatives, costs only. A word in the title is not a method. One paper in the pack does exactly this."),
 ("Cost-minimisation analysis is legitimate only when:",
  ["the new option is the cheaper one", "equivalent outcomes have been demonstrated, not assumed",
   "all costs are in rupees", "a decision model was used"], 1,
  "If outcomes differ, ignoring them is not minimising cost; it is ignoring the question."),
]),
("Session 3", "Framing the question", [
 ("The Indian Reference Case asks for which base-case perspective?",
  ["Provider", "Payer only", "Abridged societal", "Full societal, including productivity"], 2,
  "Payer and patient direct costs, including out-of-pocket. Payer perspective reported separately."),
 ("The comparator in an Indian evaluation should normally be:",
  ["doing nothing", "current practice", "the best practice anywhere in the world", "the cheapest option available"], 1,
  "Do-nothing is right only when it is genuinely what happens now. A weak comparator flatters every result downstream."),
 ("HPV vaccination of 14-year-olds evaluated over a one-year horizon would look:",
  ["cost-effective", "like a waste of money", "unaffected by the horizon", "cost-saving"], 1,
  "All the cost falls in year one; the cancers prevented come decades later. The horizon decides the answer."),
]),
("Session 4", "Costing your own service", [
 ("A ₹5,00,000 machine will last ten years. What should one year of a costing study carry?",
  ["₹5,00,000", "₹50,000", "about ₹58,600, annuitised at 3%", "Nothing: it has already been bought"], 2,
  "5,00,000 ÷ 8.530. Straight-line ₹50,000 ignores discounting; zero ignores opportunity cost."),
 ("A PM-JAY package rate is best described as:",
  ["the cost of the procedure", "a price", "an out-of-pocket payment", "an ICER"], 1,
  "What the payer pays. In 2018 about 42% of packages were priced below half their measured cost."),
 ("Costing a service top down, dividing the department's spend by its patients, tends to miss:",
  ["overheads", "variation between patients", "salaries", "the department's budget"], 1,
  "Top down sees every rupee but makes every patient average. Bottom up sees variation but misses what nobody counted."),
]),
("Session 5", "Outcomes, discounting and the ICER", [
 ("Two years lived at a utility of 0.5 are worth how many QALYs?",
  ["0.5", "1", "2", "2.5"], 1, "0.5 × 2. The QALY's power and its controversy in one line."),
 ("A programme costs ₹3,400 more than the comparator and gains 0.13 QALYs. What is the ICER?",
  ["₹204 per QALY", "₹26,154 per QALY", "₹4,420 per QALY", "₹442 per QALY"], 1,
  "3,400 ÷ 0.13. ₹204 is the average ratio, total cost over total QALYs, meaningless for a decision."),
 ("India's discount rate for costs and outcomes is:",
  ["0%", "3%", "5%", "10%"], 1, "3%, varied from 0 to 5% in sensitivity analysis, with undiscounted results also reported."),
]),
("Session 6", "Thresholds, opportunity cost and budget impact", [
 ("The rule that cost per DALY below one to three times GDP per capita is cost-effective is:",
  ["India's official threshold", "WHO's current recommendation",
   "a heuristic that its own authors have argued should be retired", "derived from India's health budget"], 2,
  "Demand-side, not linked to any budget. Two papers in the pack still use it; ask whether it changes their verdict."),
 ("India's recent estimate of willingness to pay per QALY is closest to:",
  ["₹54,881", "₹2,12,307", "₹20,000", "₹10,00,000"], 1,
  "Chugh et al. 2026. ₹54,881 is the ICER for VIA screening from Session 2, a distractor on purpose."),
 ("An intervention is cost-effective. Which question does that leave unanswered?",
  ["Is it good value per QALY?", "Can we afford it for everyone this year?",
   "Is its ICER below the threshold?", "Does it produce health?"], 1,
  "Budget impact is a separate test. Cost-effective is not the same as funded."),
]),
("Session 7", "Modelling in half an hour", [
 ("Drug A costs ₹10,000 and cures 80%. Those not cured need second-line treatment costing ₹20,000. Expected cost per patient?",
  ["₹10,000", "₹12,000", "₹14,000", "₹30,000"], 2, "10,000 + 0.2 × 20,000. Roll back: multiply along each path, then add."),
 ("The Markov assumption means:",
  ["nobody can move back to a healthier state", "the next move depends only on the current state, not on history",
   "every cycle is one year", "every patient has the same costs"], 1,
  "Memoryless. If time since surgery changes the risk, the model must be built to remember it."),
]),
("Session 8", "Handling uncertainty", [
 ("Probabilistic sensitivity analysis deals with:",
  ["parameter uncertainty", "structural uncertainty", "heterogeneity between patients", "all three"], 0,
  "Only parameters. If the model is built wrong, all ten thousand simulations are wrong together."),
 ("Which distribution suits a cost parameter in a PSA?",
  ["Normal", "Gamma or log-normal", "Beta", "It does not matter"], 1,
  "Costs cannot be negative and are right-skewed. A normal distribution lets the model draw a negative cost."),
 ("\"An 80% probability of being cost-effective\" means:",
  ["80% of patients will benefit", "the intervention works in 80% of trials",
   "in 80% of simulations it was good value at that threshold, given the model", "we can be sure it is cost-effective"], 2,
  "And a one-in-five chance of being wrong. 'Probably', not 'yes'."),
]),
("Close", "Close", [
 ("How confident are you now that you could appraise a published economic evaluation in your field?",
  ["Not at all", "A little", "Fairly", "Very"], None, "Compare with the Session 1 baseline. Opinion; no right answer."),
]),
]

CSS = """
.wrap{max-width:1000px}
.q{border:1px solid var(--line);border-left:4px solid var(--band);padding:8px 14px;margin:10px 0;break-inside:avoid}
.q .n{font-weight:700;color:var(--band);margin-right:6px}
.q ol{margin:6px 0 4px 1.4em;padding:0} .q li{margin:2px 0}
.q li.ok{font-weight:700;color:var(--good)}
.q .why{font-size:.9rem;color:var(--ink-dim);margin:4px 0 0}
.q .op{font-size:.8rem;font-weight:700;color:var(--warn)}
@media print{ .q{padding:5px 10px;margin:6px 0} body{font-size:9.5pt} }
"""

def html_page():
    out = [f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Poll question bank</title>
<meta name="description" content="Phone-poll questions for each session of the faculty workshop, with answers and explanations.">
<link rel="stylesheet" href="../assets/css/tokens.css"><link rel="stylesheet" href="../assets/css/handout.css">
<style>{CSS}</style></head><body>
<header class="band"><div class="in"><p class="kick">Facilitators &nbsp;·&nbsp; Economic Evaluation in Health &nbsp;·&nbsp; 23 September 2026</p>
<h1>Poll question bank</h1></div></header>
<div class="wrap">
<p class="intro">Two or three questions at the end of each session, answered on phones. Correct answers are
in <strong style="color:var(--good)">green</strong>. Launch the question, give about 40 seconds, show the
distribution, then read the explanation. Where the room splits, that is the discussion; do not rush it.</p>
<p class="small dim">Every number here already appears on a slide. Two questions are opinion questions with
no right answer; they give a before-and-after on confidence.</p>"""]
    k = 0
    for t, sess, qs in BANK:
        out.append(f"<h2>{t} &nbsp;·&nbsp; {E(sess)}</h2>")
        for q, opts, ok, why in qs:
            k += 1
            out.append(f'<div class="q"><p style="margin:0"><span class="n">Q{k}</span>{E(q)}'
                       + (' <span class="op">opinion</span>' if ok is None else '') + '</p><ol type="A">')
            for i, o in enumerate(opts):
                out.append(f'<li{" class=ok" if i == ok else ""}>{E(o)}</li>')
            out.append(f'</ol><p class="why">{E(why)}</p></div>')
    out.append("</div></body></html>")
    return "\n".join(out), k

def xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    wb = Workbook(); ws = wb.active; ws.title = "Questions"
    hdr = ["#", "Session", "Topic", "Question type", "Question", "Answer A", "Answer B",
           "Answer C", "Answer D", "Correct answer", "Explanation"]
    ws.append(hdr)
    for c in ws[1]: c.font = Font(bold=True)
    k = 0
    for t, sess, qs in BANK:
        for q, opts, ok, why in qs:
            k += 1
            ws.append([k, t, sess, "Multiple choice", q, *opts,
                       ("" if ok is None else "ABCD"[ok]), why])
    widths = [4, 9, 30, 15, 60, 28, 28, 28, 28, 10, 60]
    for i, w in enumerate(widths):
        ws.column_dimensions["ABCDEFGHIJK"[i]].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    n = wb.create_sheet("How to import")
    for line in [
        "Socrative imports quizzes only from its own Excel template.",
        "Download the current import template from Socrative (its help pages describe where; the menu has moved between versions).",
        "Paste columns E to J from the Questions sheet into the template's question rows,",
        "one quiz per session (use the Session column to split), then import.",
        "Opinion questions (Q1, Q23) have no correct answer: leave that cell blank.",
    ]:
        n.append([line])
    n.column_dimensions["A"].width = 100
    wb.save(HERE / "poll-bank.xlsx")

if __name__ == "__main__":
    page, k = html_page()
    (HERE / "poll-bank.html").write_text(page)
    xlsx()
    print(k, "questions; em dashes:", page.count("—"))
