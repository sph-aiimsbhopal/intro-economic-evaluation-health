#!/usr/bin/env python3
"""Turn the extracted slide geometry into an editable PowerPoint deck.

    python3 build_pptx.py <deck> [<indir>] [<outdir>]

Every text block becomes a real PowerPoint text box at the same position and
size as in the HTML deck, so the wording can be edited. Only SVG diagrams,
equations and embedded widgets arrive as pictures. Speaker notes are carried
into the notes pane.
"""
import json, pathlib, sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

deck = sys.argv[1]
indir = pathlib.Path(sys.argv[2] if len(sys.argv) > 2 else "/home/claude/pptx")
outdir = pathlib.Path(sys.argv[3] if len(sys.argv) > 3 else "/home/claude/pptx/out")
outdir.mkdir(parents=True, exist_ok=True)

PX = 9525                      # EMU per CSS px on a 1280 px wide, 13.333 in slide
FONT = "Arial"
ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT,
         "justify": PP_ALIGN.JUSTIFY, "start": PP_ALIGN.LEFT}

def emu(v): return Emu(int(round(v * PX)))
def pt(px): return Pt(round(px * 0.75, 1))
def col(h): return RGBColor.from_string(h) if h else None

def set_bullet(para, char):
    """python-pptx has no bullet API; write the paragraph properties directly."""
    pPr = para._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    if char is None:
        pPr.append(pPr.makeelement(qn("a:buNone"), {}))
        return
    pPr.set("marL", "285750"); pPr.set("indent", "-285750")
    if char.endswith("."):
        attrs = {"type": "arabicPeriod"}
        n = char[:-1]
        if n.isdigit() and int(n) > 1:
            attrs["startAt"] = n          # keep numbering across a split column
        e = pPr.makeelement(qn("a:buAutoNum"), attrs)
    else:
        f = pPr.makeelement(qn("a:buFont"), {"typeface": FONT})
        pPr.append(f)
        e = pPr.makeelement(qn("a:buChar"), {"char": char})
    pPr.append(e)

def add_text(slide, it):
    r = it["rect"]
    box = slide.shapes.add_textbox(emu(r["x"]), emu(r["y"]) - Emu(20000),
                                   emu(r["w"]), emu(max(r["h"], 12)))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if it.get("vcenter"):
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        box.top, box.height = emu(r["y"]), emu(r["h"])
    for i, para in enumerate(it["paras"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN.get(it.get("align", "left"), PP_ALIGN.LEFT)
        p.line_spacing = round(it.get("line", 1.3) / 1.2, 3)   # PowerPoint multiples are of 1.2 lines
        p.space_after = Pt(4 if len(it["paras"]) > 1 else 0)
        set_bullet(p, para.get("bullet"))
        for run in para["runs"]:
            r_ = p.add_run()
            r_.text = run["t"]
            f = r_.font
            f.name = FONT
            f.size = pt(run.get("sz") or it["sz"])
            f.bold = bool(run.get("b"))
            f.italic = bool(run.get("i"))
            c = col(run.get("color") or it.get("color"))
            if c is not None:
                f.color.rgb = c
            if run.get("spc"):
                r_.font._rPr.set("spc", str(int(round(run["spc"] * 75))))   # 1/100 pt
    return box

def add_rect(slide, it):
    r = it["rect"]
    shape_kind = MSO_SHAPE.ROUNDED_RECTANGLE if (it.get("radius") or 0) >= 4 else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(shape_kind, emu(r["x"]), emu(r["y"]), emu(r["w"]), emu(r["h"]))
    if shape_kind == MSO_SHAPE.ROUNDED_RECTANGLE:      # keep the corner subtle
        sh.adjustments[0] = 0.04
    if it.get("fill"):
        sh.fill.solid(); sh.fill.fore_color.rgb = col(it["fill"])
    else:
        sh.fill.background()
    if it.get("line"):
        sh.line.color.rgb = col(it["line"]); sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    if it.get("leftBar"):
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(r["x"]), emu(r["y"]),
                                     emu(it["leftBar"]["w"]), emu(r["h"]))
        bar.fill.solid(); bar.fill.fore_color.rgb = col(it["leftBar"]["color"])
        bar.line.fill.background(); bar.shadow.inherit = False
    return sh

def add_table(slide, it):
    r, rows = it["rect"], it["rows"]
    ncol = max(len(x) for x in rows)
    g = slide.shapes.add_table(len(rows), ncol, emu(r["x"]), emu(r["y"]),
                               emu(r["w"]), emu(r["h"])).table
    g.first_row = rows[0][0]["head"] if rows and rows[0] else False
    g.horz_banding = False
    widths = it.get("widths") or []
    for i in range(ncol):
        if i < len(widths) and widths[i]:
            g.columns[i].width = emu(widths[i])
    for ri, row in enumerate(rows):
        for ci in range(ncol):
            cell = g.cell(ri, ci)
            cell.margin_left = cell.margin_right = Emu(60000)
            cell.margin_top = cell.margin_bottom = Emu(35000)
            cell.vertical_anchor = MSO_ANCHOR.TOP
            src = row[ci] if ci < len(row) else None
            tf = cell.text_frame
            tf.word_wrap = True
            if not src:
                cell.fill.background(); continue
            if src.get("fill"):
                cell.fill.solid(); cell.fill.fore_color.rgb = col(src["fill"])
            else:
                cell.fill.background()
            for i, para in enumerate(src["paras"] or [{"runs": [{"t": ""}]}]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = ALIGN.get(src.get("align", "left"), PP_ALIGN.LEFT)
                p.line_spacing = 1.04
                for run in para["runs"]:
                    r_ = p.add_run(); r_.text = run["t"]
                    f = r_.font
                    f.name = FONT; f.size = pt(it["sz"])
                    f.bold = bool(run.get("b")) or src["head"]
                    f.italic = bool(run.get("i"))
                    c = col(run.get("color") or src.get("color"))
                    if c is not None:
                        f.color.rgb = c
    return g

def build(deck):
    data = json.loads((indir / f"{deck}.json").read_text())
    prs = Presentation()
    prs.slide_width, prs.slide_height = emu(1280), emu(720)
    blank = prs.slide_layouts[6]
    for s in data["slides"]:
        slide = prs.slides.add_slide(blank)
        for it in s["items"]:
            if it["kind"] == "rect":
                add_rect(slide, it)
            elif it["kind"] == "text":
                add_text(slide, it)
            elif it["kind"] == "table":
                add_table(slide, it)
            elif it["kind"] == "image" and it.get("file"):
                r = it["rect"]
                slide.shapes.add_picture(it["file"], emu(r["x"]), emu(r["y"]),
                                         emu(r["w"]), emu(r["h"]))
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
    path = outdir / f"{deck}.pptx"
    prs.save(path)
    print(f"{deck}: {len(data['slides'])} slides -> {path} "
          f"({path.stat().st_size/1024:.0f} KB)")

build(deck)
